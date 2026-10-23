import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title
p_title = doc.add_paragraph()
r_title = p_title.add_run("Project Master Blueprint & Status Report")
r_title.bold = True
r_title.font.size = Pt(20)
r_title.font.color.rgb = RGBColor(24, 43, 73)

p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(14)
r_meta = p_meta.add_run("Project: SLM vs. LLM Phishing Detection | Date: October 2026\nAuthor: Satvik Patel | Mentor: Prof. Amit Sir | Current Milestone: End of Week 5")
r_meta.font.size = Pt(10)
r_meta.font.color.rgb = RGBColor(100, 100, 100)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    for r in h.runs:
        r.font.color.rgb = RGBColor(24, 43, 73)
    return h

# 1. The Big Picture
add_heading("1. The Big Picture: What is This Project About?", level=1)
doc.add_paragraph(
    "Every day, cybercriminals use generative AI (like ChatGPT) to write deceptive phishing emails that look completely professional and contain no traditional spam words. "
    "Traditional spam filters fail because they rely on fixed keywords. Large cloud models like GPT-4 can read emails, but sending enterprise emails to cloud APIs is expensive and poses serious privacy and compliance risks."
)
doc.add_paragraph(
    "Our Core Goal: Can a compact Small Language Model (SLM) with only 1.5 billion parameters running privately on a standard computer or graphics card match or beat massive cloud models while eliminating false alarms?"
)

# 2. Master Roadmap
add_heading("2. Master Project Blueprint (8-Week Roadmap)", level=1)
steps = [
    ("Week 1", "Literature Review & Experimental Setup", "COMPLETED", "Defined research hypotheses and protocols."),
    ("Week 2", "Data Acquisition & Vault Splitting", "COMPLETED", "Assembled 29,544 emails; locked test set (5,909 emails)."),
    ("Week 3", "Classical Machine Learning Baselines", "COMPLETED", "TF-IDF + Logistic Regression achieved 97.36% Accuracy."),
    ("Week 4", "Zero-Shot AI Baseline Testing", "COMPLETED", "Discovered untuned models fail (33% to 72% false alarm rate)."),
    ("Week 5", "Arm A: Direct Fine-Tuning (RTX 5070)", "COMPLETED", "Fine-tuned Qwen2.5-1.5B: Achieved 99.15% Accuracy and 0.9900 F1!"),
    ("Week 6", "Arm B: Reasoning (CoT) Fine-Tuning", "UP NEXT", "Teach the model to explain WHY an email is phishing before deciding."),
    ("Week 7", "AI Adversarial Attack Stress Testing", "UPCOMING", "Test models against ChatGPT-rephrased evasive phishing emails."),
    ("Week 8", "Final Test Set Unlock & Thesis Paper", "UPCOMING", "Unlock the 5,909 test emails for final grading and paper submission.")
]

t1 = doc.add_table(rows=len(steps)+1, cols=4)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Phase", "Task Description", "Status", "Key Objective / Milestone"]
for j, h in enumerate(headers):
    t1.rows[0].cells[j].text = h
for i, row in enumerate(steps):
    for j, val in enumerate(row):
        t1.rows[i+1].cells[j].text = val

for row in t1.rows:
    for c in row.cells:
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

# 3. What Work is Done
add_heading("3. What Work is Already Completed (Weeks 1 to 5)", level=1)
doc.add_paragraph(
    "1. Data Preparation: 20,681 train emails, 2,954 validation emails, and 5,909 test emails locked in a cryptographic vault.\n"
    "2. Classical Baselines: Built and verified TF-IDF + Logistic Regression (97.36% accuracy, 2.17% False Positive Rate).\n"
    "3. Zero-Shot Revelation: Proved that untrained models (both 1.5B and 27B) are unusable in production due to massive false alarms (33% to 72% FPR).\n"
    "4. Arm A Fine-Tuning (Huge Success): Fine-tuned Qwen2.5-1.5B for 3 epochs on an NVIDIA RTX 5070 GPU. Reached 99.15% Accuracy, 0.9900 F1-score, and cut false positives down to 0.88%!"
)

# 4. Scorecard
add_heading("4. Current Model Scorecard (Validation Set: 2,954 Emails)", level=1)
scorecard = [
    ("Zero-Shot Big AI (27B)", "32.54%", "39.26%", "72.44% (Very High)", "0.3312"),
    ("Zero-Shot Small AI (1.5B)", "63.64%", "59.70%", "33.47% (High)", "0.5811"),
    ("Classical ML (TF-IDF + Logistic)", "97.36%", "96.71%", "2.17% (Good)", "0.9687"),
    ("Our Fine-Tuned Small AI (Arm A)", "99.15%", "99.20%", "0.88% (Outstanding)", "0.9900")
]

t2 = doc.add_table(rows=len(scorecard)+1, cols=5)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
headers2 = ["Model / Approach", "Accuracy", "Caught Threats (Recall)", "False Alarms (FPR)", "Overall Score (F1)"]
for j, h in enumerate(headers2):
    t2.rows[0].cells[j].text = h
for i, row in enumerate(scorecard):
    for j, val in enumerate(row):
        t2.rows[i+1].cells[j].text = val

for row in t2.rows:
    for c in row.cells:
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

# 5. What Work is Left
add_heading("5. What Work is Left To Do (Weeks 6 to 8)", level=1)
doc.add_paragraph(
    "Phase 1: Week 6 — Arm B (Reasoning Distillation)\n"
    "- Goal: Teach the model to generate reasoning traces before deciding (e.g., analyzing domain mismatches and urgency).\n"
    "- Execution: Distill explanations from a teacher model and fine-tune on the RTX 5070 (~5.5 to 6.5 hours).\n\n"
    "Phase 2: Week 7 — Adversarial Stress Testing\n"
    "- Goal: Test if Arm B is truly superior to Arm A when attacked by sophisticated AI.\n"
    "- Test: Rephrase phishing emails with ChatGPT to remove obvious spam keywords and evaluate whether Arm A gets fooled while Arm B catches the subtle deception.\n\n"
    "Phase 3: Week 8 — Final Evaluation & Research Paper\n"
    "- Goal: Unlock the 5,909 test emails for the final, official benchmark.\n"
    "- Deliverables: Speed and cost trade-off charts, complete project presentation, and the final thesis report for Prof. Amit Sir."
)

output_file = r"c:\Users\satvi\Desktop\Amit Sir\Project_Master_Blueprint_And_Status.docx"
doc.save(output_file)
print("Successfully generated:", output_file)
